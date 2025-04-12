import os
import pandas as pd
from docx import Document

# 尋找檔案


def find_file_by_ext(extension):
    for file in os.listdir():
        if file.lower().endswith(extension):
            return file
    return None


excel_file = find_file_by_ext(".xlsx")
word_template = find_file_by_ext(".docx")

if not excel_file or not word_template:
    print("找不到 Excel 或 Word 檔案")
    exit()

df = pd.read_excel(excel_file)
output_dir = "table"
os.makedirs(output_dir, exist_ok=True)

# 替換 run 內的佔位符，保留樣式


def replace_placeholders_in_runs(paragraph, row_data):
    for run in paragraph.runs:
        for key, value in row_data.items():
            placeholder = "{{" + key + "}}"
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, str(value))


# 開始產生 Word 檔
for idx, row in df.iterrows():
    doc = Document(word_template)

    for para in doc.paragraphs:
        replace_placeholders_in_runs(para, row)

    output_path = os.path.join(output_dir, f"table_{idx}.docx")
    doc.save(output_path)

print(f"✅ 成功產生 {len(df)} 份文件，保留樣式 ✅")
